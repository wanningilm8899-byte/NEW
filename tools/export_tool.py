from __future__ import annotations
from datetime import datetime
from docx import Document
from docx.shared import Pt
from core.paths import OUTPUT_DIR
from agent.risk_brief import baseline_brief, scenario_brief, data_audit_brief


def _clean(s: str) -> str:
    return (s or '').replace('**','').replace('`','').strip()

def _add_markdownish(doc: Document, text: str):
    for raw in (text or '').splitlines():
        line=raw.strip()
        if not line: continue
        if line.startswith('### '): doc.add_heading(_clean(line[4:]), level=2)
        elif line.startswith('- '): doc.add_paragraph(_clean(line[2:]), style='List Bullet')
        else: doc.add_paragraph(_clean(line))


def export_risk_report(baseline: dict, scenario: dict|None, validation: dict|None,
                       resolved_request: dict|None, snapshot_id: str='', role: str='risk') -> str:
    doc=Document()
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.add_heading('Water Risk Copilot｜上游供应链水风险智能决策简报',0)
    doc.add_paragraph(f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    if snapshot_id: doc.add_paragraph(f'Input Snapshot：{snapshot_id}')
    if baseline:
        doc.add_paragraph(f"Run ID：{baseline.get('run_id','—')}  |  Input Hash：{baseline.get('input_hash','—')}")

    doc.add_heading('1. 数据输入与质量状态',1)
    _add_markdownish(doc, data_audit_brief(validation, resolved_request))

    doc.add_heading('2. Baseline 结果与风险驱动',1)
    _add_markdownish(doc, baseline_brief(baseline, role=role))
    if baseline and baseline.get('materials'):
        t=doc.add_table(rows=1,cols=7); t.style='Table Grid'
        hdr=['材料','状态','PRWI','采购覆盖','可评分覆盖','Unknown','证据可信度']
        for i,x in enumerate(hdr): t.rows[0].cells[i].text=x
        for m in baseline['materials']:
            vals=[m.get('material'),m.get('status'),m.get('PRWI'),m.get('procurement_coverage'),m.get('scored_coverage'),m.get('unknown_share'),m.get('overall_confidence')]
            row=t.add_row().cells
            for i,v in enumerate(vals): row[i].text='—' if v is None else str(v)

    doc.add_heading('3. Scenario 压力测试',1)
    _add_markdownish(doc, scenario_brief(scenario,baseline,role) if scenario else '本次未运行 Scenario。')

    doc.add_heading('4. 数据缺口、Proxy 与假设',1)
    if baseline:
        for label,key in [('Missing','missing_fields'),('Proxy','proxy_fields'),('Warnings','warnings'),('Assumptions','assumptions')]:
            vals=baseline.get(key) or []
            doc.add_paragraph(f'{label}: '+('无' if not vals else str(vals)))
    if resolved_request:
        doc.add_paragraph('数据版本：'+str(resolved_request.get('data_version')))

    doc.add_heading('5. 版本与可追溯信息',1)
    if baseline:
        for k in ['engine_version','formula_version','parameter_version','data_version','calculated_at','agent_adapter_version']:
            doc.add_paragraph(f'{k}: {baseline.get(k,"—")}')
    if scenario:
        doc.add_paragraph('scenario_version: '+str(scenario.get('scenario_version','—')))
        doc.add_paragraph('baseline_run_id: '+str(scenario.get('baseline_run_id','—')))
        doc.add_paragraph('baseline_input_hash: '+str(scenario.get('baseline_input_hash','—')))

    doc.add_heading('6. 解释边界',1)
    doc.add_paragraph('PRWI 用于相对风险筛查，不等同于损失概率或财务损失预测。AI 不生成或修改风险数字；真实采购调整、供应商替换和投资决策需人工确认。')

    out=OUTPUT_DIR/f'Water_Risk_Copilot_Brief_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
    doc.save(out)
    return str(out)
